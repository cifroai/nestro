# -*- coding: utf-8 -*-
"""
Расчёт показателей качества цементирования по реестру РВП.

Вход:  data/Реестр_2026.xlsx (лист «Скважины», как прислан заказчиком; строки с 5-й)
Выход: data/metrics.json            — числа для слайдов
       output/Расчёт_цементирование_2026.xlsx — таблица расчётов с формулами (проверка каждого числа)

Правила (см. README в книге расчётов):
* база анализа — скважины, законченные бурением в 2026 г. (дата окончания заполнена и год = 2026);
  записи без даты окончания показаны отдельно и в расчёте чувствительности;
* единица оценки — колонна (операция цементирования); ОС и весь интервал — раздельно;
* «н/ц» — колонна не цементировалась → исключается из оценки качества;
  «Не писали», «нет инф», пусто → «н/д» (не восстанавливается);
* ВПЦ: пара = числовые план и факт в одной колонке-паре реестра; отклонение = факт − план (> 0 — недоподъём);
* подтверждённых оснований технической невалидности в реестре нет (столбец «Примечание» пуст) → исключений нет.
"""
import datetime as dt
import json
import os
import statistics as st

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.path.join(ROOT, "data", "Реестр_2026.xlsx")
OUT_JSON = os.path.join(ROOT, "data", "metrics.json")
OUT_XLSX = os.path.join(ROOT, "output", "Расчёт_цементирование_2026.xlsx")

STRINGS = [  # ключ, название, колонки реестра (ОС, весь, план, факт), диаметр
    ("tech", "Техническая", ("N", "O", "P", "Q"), "245"),
    ("prod", "Эксплуатационная", ("S", "T", "U", "V"), None),  # диаметр — столбец R
    ("liner", "Хвостовик", ("W", "X", "Y", "Z"), "114"),
]
SNAME = {k: n for k, n, _, _ in STRINGS}
STATUS_TXT = {"н/ц": "не цементировался", "не писали": "АКЦ не записан", "нет инф": "нет информации"}
YEAR = 2026


def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def status(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return "пусто"
    if num(v) is not None:
        return "число"
    return STATUS_TXT.get(str(v).strip().lower(), "текст: " + str(v).strip())


def pdate(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, str) and v.strip():
        return dt.datetime.strptime(v.strip(), "%d.%m.%y").date()
    return None


def stats(vals):
    vals = [v for v in vals if v is not None]
    if not vals:
        return {"n": 0, "mean": None, "median": None, "min": None, "max": None}
    return {"n": len(vals), "mean": sum(vals) / len(vals), "median": st.median(vals), "min": min(vals), "max": max(vals)}


# ---------------------------------------------------------------- чтение
def read_registry():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb["Скважины"]
    wells, ops, raw_rows = [], [], []
    for r in range(5, ws.max_row + 1):
        wid = ws[f"D{r}"].value
        if wid is None or not str(wid).strip():
            continue
        raw_rows.append(r)
        wid = str(wid).strip()
        field = str(ws[f"B{r}"].value).strip()
        pad = ws[f"C{r}"].value
        pad = str(pad).strip() if pad is not None and str(pad).strip() else None
        up = wid.upper()
        if "ГС" in up:
            prof, prof_src = "ГС", "суффикс «ГС» в номере"
        elif "ННС" in up:
            prof, prof_src = "НН", "суффикс «ННС» в номере"
        else:
            prof, prof_src = "НН", "нет суффикса «ГС» (принято НН, подтвердить)"
        start, end = pdate(ws[f"L{r}"].value), pdate(ws[f"M{r}"].value)
        w = dict(row=r, well=wid, field=field, pad=pad, profile=prof, profile_src=prof_src,
                 contractor=ws[f"E{r}"].value, start=start, end=end,
                 end_year=end.year if end else None,
                 quarter=f"{(end.month - 1) // 3 + 1} кв. {end.year}" if end else None,
                 base=1 if end and end.year == YEAR else 0,
                 note=ws[f"AA{r}"].value)
        wells.append(w)
        for key, name, cols, diam in STRINGS:
            vals = [ws[f"{c}{r}"].value for c in cols]
            d = diam or (str(ws[f"R{r}"].value).strip() if ws[f"R{r}"].value is not None else None)
            sts = [status(v) for v in vals]
            cemented = 0 if "не цементировался" in sts else 1
            ops.append(dict(w, key=key, string=name, diam=d, cemented=cemented,
                            os_raw=vals[0], full_raw=vals[1], plan_raw=vals[2], fact_raw=vals[3],
                            os=num(vals[0]) if cemented else None, full=num(vals[1]) if cemented else None,
                            plan=num(vals[2]) if cemented else None, fact=num(vals[3]) if cemented else None,
                            akc_status=sts[0] if sts[0] == sts[1] else f"ОС: {sts[0]}; весь: {sts[1]}",
                            src_cells=f"{cols[0]}{r}:{cols[3]}{r}"))
    for o in ops:
        o["pair"] = 1 if (o["plan"] is not None and o["fact"] is not None) else 0
        o["dev"] = (o["fact"] - o["plan"]) if o["pair"] else None
    return wells, ops, raw_rows, wb


# ---------------------------------------------------------------- метрики
def compute(wells, ops):
    B = [o for o in ops if o["base"] and o["cemented"]]
    M = {"year": YEAR, "has_2025": False,
         "note_2025": "Построчные данные и опубликованные агрегаты за 2025 г. не предоставлены — сравнение 2025 → 2026 невозможно."}
    bw = [w for w in wells if w["base"]]
    M["coverage"] = {
        "registry_wells": len(wells),
        "wells": len(bw),
        "wells_nodate": len([w for w in wells if not w["end"]]),
        "wells_nn": len([w for w in bw if w["profile"] == "НН"]),
        "wells_gs": len([w for w in bw if w["profile"] == "ГС"]),
        "wells_nn_suffix": len([w for w in bw if w["profile"] == "НН" and "ННС" in w["well"].upper()]),
        "strings_total": len([o for o in ops if o["base"]]),
        "strings_uncemented": len([o for o in ops if o["base"] and not o["cemented"]]),
        "strings_cemented": len(B),
        "strings_akc_any": len([o for o in B if o["os"] is not None or o["full"] is not None]),
        "strings_akc_oh": len([o for o in B if o["os"] is not None]),
        "strings_akc_full": len([o for o in B if o["full"] is not None]),
        "strings_akc_missing": len([o for o in B if o["os"] is None and o["full"] is None]),
        "toc_pairs": sum(o["pair"] for o in B),
        "wells_with_akc": len({o["well"] for o in B if o["full"] is not None or o["os"] is not None}),
    }
    # статусы отсутствующих данных
    miss = {}
    for o in B:
        if o["os"] is None and o["full"] is None:
            miss[o["akc_status"]] = miss.get(o["akc_status"], 0) + 1
    M["akc_missing_by_status"] = miss

    def grp(filt):
        sel = [o for o in B if filt(o)]
        return {"n_strings": len(sel), "bond_oh": stats([o["os"] for o in sel]), "bond_full": stats([o["full"] for o in sel]),
                "toc": {**stats([o["dev"] for o in sel]), "plan_median": (st.median([o["plan"] for o in sel if o["pair"]]) if any(o["pair"] for o in sel) else None),
                        "fact_median": (st.median([o["fact"] for o in sel if o["pair"]]) if any(o["pair"] for o in sel) else None),
                        "under_n": len([o for o in sel if o["pair"] and o["dev"] > 0])}}

    M["by_string"] = {k: grp(lambda o, k=k: o["key"] == k) for k, *_ in STRINGS}
    M["all_strings"] = grp(lambda o: True)
    M["by_type"] = {p: {k: grp(lambda o, p=p, k=k: o["profile"] == p and o["key"] == k) for k, *_ in STRINGS} for p in ("НН", "ГС")}
    fields = sorted({w["field"] for w in wells}, key=lambda f: -len([w for w in bw if w["field"] == f]))
    M["fields"] = fields
    M["by_field"] = {f: {"wells": len([w for w in bw if w["field"] == f]),
                         "wells_nn": len([w for w in bw if w["field"] == f and w["profile"] == "НН"]),
                         "wells_gs": len([w for w in bw if w["field"] == f and w["profile"] == "ГС"]),
                         "strings_cemented": len([o for o in B if o["field"] == f]),
                         "strings_akc_full": len([o for o in B if o["field"] == f and o["full"] is not None]),
                         "toc_pairs": sum(o["pair"] for o in B if o["field"] == f),
                         **{k: grp(lambda o, f=f, k=k: o["field"] == f and o["key"] == k) for k, *_ in STRINGS},
                         "all": grp(lambda o, f=f: o["field"] == f)} for f in fields}
    M["by_field_type"] = {f: {p: {k: grp(lambda o, f=f, p=p, k=k: o["field"] == f and o["profile"] == p and o["key"] == k)
                                  for k, *_ in STRINGS} for p in ("НН", "ГС")} for f in fields}
    contractors = sorted({w["contractor"] for w in wells if w["contractor"]})
    M["contractors"] = contractors
    M["by_contractor"] = {c: {"wells": len([w for w in bw if w["contractor"] == c]),
                              **{k: grp(lambda o, c=c, k=k: o["contractor"] == c and o["key"] == k) for k, *_ in STRINGS}} for c in contractors}
    quarters = sorted({o["quarter"] for o in B if o["quarter"]})
    M["quarters"] = quarters
    M["by_quarter"] = {q: {"wells": len([w for w in bw if w["quarter"] == q]),
                           **{k: grp(lambda o, q=q, k=k: o["quarter"] == q and o["key"] == k) for k, *_ in STRINGS}} for q in quarters}
    # пороги — только справочно (критерий не подтверждён)
    M["thresholds"] = {}
    for k, *_ in STRINGS:
        for key in ("os", "full"):
            vals = [o[key] for o in B if o["key"] == k and o[key] is not None]
            M["thresholds"][f"{k}_{key}"] = {"n": len(vals), "ge80": len([v for v in vals if v >= 0.80]), "ge90": len([v for v in vals if v >= 0.90])}
    # точки для диаграмм
    M["points"] = [{"well": o["well"], "field": o["field"], "profile": o["profile"], "key": o["key"], "string": o["string"], "diam": o["diam"],
                    "os": o["os"], "full": o["full"], "plan": o["plan"], "fact": o["fact"], "dev": o["dev"], "quarter": o["quarter"],
                    "pad": o["pad"], "contractor": o["contractor"]} for o in B]
    # наихудшие операции
    valid_full = [o for o in B if o["full"] is not None]
    M["worst_bond"] = [dict(well=o["well"], pad=o["pad"], field=o["field"], profile=o["profile"], string=o["string"], diam=o["diam"],
                            os=o["os"], full=o["full"], dev=o["dev"], contractor=o["contractor"])
                       for o in sorted(valid_full, key=lambda o: (o["full"], o["os"] if o["os"] is not None else 9))[:7]]
    under = [o for o in B if o["pair"] and o["dev"] > 0]
    M["worst_toc"] = [dict(well=o["well"], pad=o["pad"], field=o["field"], profile=o["profile"], string=o["string"], diam=o["diam"],
                           plan=o["plan"], fact=o["fact"], dev=o["dev"], full=o["full"], contractor=o["contractor"])
                      for o in sorted(under, key=lambda o: -o["dev"])[:7]]
    # чувствительность: исключений по невалидности нет; вариант «+ записи без даты окончания»
    Ball = [o for o in ops if o["cemented"]]
    M["sensitivity"] = {
        "invalid_confirmed": 0,
        "base": {k: {"full": stats([o["full"] for o in B if o["key"] == k]), "os": stats([o["os"] for o in B if o["key"] == k])} for k, *_ in STRINGS},
        "with_nodate": {k: {"full": stats([o["full"] for o in Ball if o["key"] == k]), "os": stats([o["os"] for o in Ball if o["key"] == k])} for k, *_ in STRINGS},
        "nodate_wells": [w["well"] for w in wells if not w["end"]],
    }
    M["missing_fields"] = ["SRTi (применение)", "Осложнения при бурении/цементировании", "Дата АКЦ",
                           "Подрядчик по цементированию", "2025 г.: построчные данные или опубликованные агрегаты",
                           "Утверждённый критерий «хорошо/плохо» (0,80 / 0,90)"]
    return M


# ---------------------------------------------------------------- проверки качества данных
def audit(wells, ops, wb):
    log = []
    ids = [w["well"] for w in wells]
    for w in wells:
        base = "".join(ch for ch in w["well"] if ch.isdigit())
        sim = [x for x in ids if x != w["well"] and "".join(ch for ch in x if ch.isdigit()) == base]
        if sim:
            log.append(("Похожие номера", w["well"], f"Есть запись {', '.join(sim)} с тем же числовым номером. Записи НЕ объединены: "
                        f"куст {w['pad'] or 'н/д'}, окончание бурения {w['end'] or 'н/д'}. Требуется подтверждение, одна ли это скважина.", f"D{w['row']}"))
        if not w["end"]:
            log.append(("Нет даты окончания бурения", w["well"], "Год окончания не подтверждён → вне базы 2026; показано в чувствительности.", f"M{w['row']}"))
        if not w["pad"]:
            log.append(("Нет куста", w["well"], "Куст не указан в реестре (н/д).", f"C{w['row']}"))
        if w["profile_src"].startswith("нет"):
            pass
    for o in ops:
        if o["base"] and o["cemented"] and o["os"] is None and o["full"] is None:
            log.append(("Нет АКЦ", o["well"], f"{o['string']}: {o['akc_status']} → н/д.", o["src_cells"]))
        if o["base"] and o["cemented"] and ((o["plan"] is None) != (o["fact"] is None)):
            log.append(("Неполная пара ВПЦ", o["well"], f"{o['string']}: план {o['plan_raw']!r}, факт {o['fact_raw']!r} → пара не сопоставляется.", o["src_cells"]))
        if o["base"] and o["cemented"] and o["full"] is not None and o["os"] is not None and o["full"] > o["os"] + 1e-9:
            log.append(("Проверить значение", o["well"], f"{o['string']}: Кобщ весь интервал ({o['full']}) > Кобщ ОС ({o['os']}).", o["src_cells"]))
        if not o["cemented"]:
            log.append(("Нецементируемая колонна", o["well"], f"{o['string']} {o['diam']} мм: «н/ц» → исключена из оценки качества.", o["src_cells"]))
    nn_nosuf = [w["well"] for w in wells if w["profile_src"].startswith("нет")]
    log.append(("Профиль без суффикса", ", ".join(nn_nosuf), "Профиль принят НН: в номере нет «ГС»; конструкция — экспл. 168/178 мм, хвостовик 114 мм "
                "цементируется или АКЦ не записан (у ГС хвостовик «н/ц»). Требует подтверждения.", "D"))
    log.append(("Невалидность замеров", "—", "Столбец «Примечание» пуст: подтверждённых оснований признать замер невалидным нет → исключений нет.", "AA5:AA27"))
    log.append(("ВПЦ «0»", "—", "Отметка 0 м трактуется как числовая глубина (до устья), как записано в реестре; текстовой отметки «устье» в реестре нет.", "P, Q, U, V"))
    return log


# ---------------------------------------------------------------- книга расчётов
HFILL = PatternFill("solid", fgColor="009C3D")
GFILL = PatternFill("solid", fgColor="00326E")
CFILL = PatternFill("solid", fgColor="E5EAF0")
F = lambda **k: Font(name="Arial", size=k.pop("size", 10), **k)
THIN = Border(bottom=Side(style="thin", color="C9D3DE"))


def hdr(ws, row, labels, col=1, fill=HFILL):
    for j, t in enumerate(labels):
        c = ws.cell(row=row, column=col + j, value=t)
        c.font = F(bold=True, color="FFFFFF", size=9)
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")


def write_book(wells, ops, M, log, raw_wb, raw_rows):
    wb = openpyxl.Workbook()
    checks = []  # (лист, ячейка, ожидаемое значение из Python) — сверка после пересчёта

    # README
    ws = wb.active
    ws.title = "README"
    lines = [
        ("Расчёт показателей качества цементирования · СП «Русвьетпетро» · 2026", True),
        ("Источник: реестр «Бурение и качество цементирования скважин РВП · 2026» (лист «Скважины»), присланный заказчиком; копия — лист «Реестр_исходный».", False),
        ("Все показатели слайдов считаются формулами на листах S1…S8 по листу «Операции». Номер листа указан в колонтитуле каждого слайда.", False),
        ("", False),
        ("Методика", True),
        ("1. База анализа — скважины, законченные бурением в 2026 г. (заполнена дата окончания бурения). Записи реестра без даты окончания (4) в базу не входят и показаны в чувствительности (S7).", False),
        ("2. Единица оценки качества — колонна (операция цементирования). Колонна типизирована по графе реестра: техническая 245 мм, эксплуатационная 168/178 мм, хвостовик 114 мм.", False),
        ("3. «н/ц» — колонна не цементировалась: исключается из оценки качества (все хвостовики ГС). «Не писали», «нет инф», пустые ячейки — «н/д», значения не восстанавливаются.", False),
        ("4. Кобщ в открытом стволе (ОС) и по всему интервалу считаются раздельно: среднее арифметическое, медиана, n, мин–макс по всей валидной выборке.", False),
        ("5. ВПЦ сопоставляется только при числовых плане и факте в одной паре граф реестра. Отклонение = факт − план, м; > 0 — недоподъём, < 0 — подъём выше плана.", False),
        ("6. Подтверждённых оснований технической невалидности замеров в реестре нет (графа «Примечание» пуста) — исключений нет.", False),
        ("7. Порог «хорошо/плохо» не утверждён: доли ≥ 0,80 и ≥ 0,90 приводятся справочно, оба.", False),
        ("8. 2025 г.: построчные данные и опубликованные агрегаты не предоставлены — статистика 2025 г. не рассчитывается.", False),
        ("9. SRTi, осложнения, даты АКЦ и подрядчик по цементированию в реестре отсутствуют — «н/д». Подрядчик в реестре — буровой.", False),
        ("10. Профиль: ГС — суффикс «ГС» в номере; остальные — НН (у одной скважины суффикс «ННС», у прочих — по отсутствию «ГС» и конструкции; требует подтверждения).", False),
        ("", False),
        ("Как проверить число со слайда: найдите лист из колонтитула слайда → строку с названием показателя. Формулы ссылаются на лист «Операции» (одна строка = одна колонна скважины).", False),
        ("Ячейки листа «Операции» с исходными значениями (графы «исх.») скопированы из реестра без изменений; числовые графы — те же значения, где они числовые.", False),
    ]
    for i, (t, b) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = F(bold=b, size=12 if i == 1 else 10, color="00326E" if b else "000000")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 150

    # Реестр_исходный — копия значений
    src = raw_wb["Скважины"]
    wr = wb.create_sheet("Реестр_исходный")
    for r in range(1, 5):
        for c in range(1, 28):
            v = src.cell(row=r, column=c).value
            if v is not None:
                wr.cell(row=r, column=c, value=v).font = F(bold=r >= 3, size=9)
    for i, r in enumerate(raw_rows):
        for c in range(1, 28):
            v = src.cell(row=r, column=c).value
            wr.cell(row=5 + i, column=c, value=v).font = F(size=9)
    wr.freeze_panes = "E5"

    # Скважины (нормализованные)
    wsw = wb.create_sheet("Скважины")
    cols = ["Скважина", "Месторождение", "Куст", "Профиль", "Основание профиля", "Буровой подрядчик", "Подрядчик по цементированию",
            "Начало бурения", "Окончание бурения", "Год окончания", "Квартал окончания", "База 2026 (1/0)", "Строка реестра"]
    hdr(wsw, 1, cols)
    for i, w in enumerate(wells, 2):
        vals = [w["well"], w["field"], w["pad"] or "н/д", w["profile"], w["profile_src"], w["contractor"], "н/д",
                w["start"], w["end"], w["end_year"] if w["end_year"] else "н/д", w["quarter"] or "н/д", w["base"], w["row"]]
        for j, v in enumerate(vals, 1):
            c = wsw.cell(row=i, column=j, value=v)
            c.font = F(size=9)
            if isinstance(v, dt.date):
                c.number_format = "DD.MM.YYYY"
    for j, wdt in enumerate([13, 22, 7, 8, 40, 11, 14, 12, 12, 9, 12, 9, 9], 1):
        wsw.column_dimensions[get_column_letter(j)].width = wdt
    wsw.freeze_panes = "B2"

    # Операции (одна строка = колонна)
    wo = wb.create_sheet("Операции")
    oc = ["Скважина", "Месторождение", "Куст", "Профиль", "Буровой подрядчик", "Квартал окончания", "База 2026 (1/0)", "Колонна",
          "Диаметр, мм", "Цементируется (1/0)", "Кобщ ОС, исх.", "Кобщ ОС", "Кобщ весь, исх.", "Кобщ весь", "ВПЦ план, исх.",
          "ВПЦ факт, исх.", "ВПЦ план, м", "ВПЦ факт, м", "Пара ВПЦ (1/0)", "Отклонение факт − план, м", "Статус АКЦ",
          "Невалидность подтверждена", "Ячейки реестра"]
    hdr(wo, 1, oc)
    N = len(ops)
    for i, o in enumerate(ops, 2):
        vals = [o["well"], o["field"], o["pad"] or "н/д", o["profile"], o["contractor"], o["quarter"] or "н/д", o["base"], o["string"],
                o["diam"], o["cemented"], o["os_raw"], o["os"], o["full_raw"], o["full"], o["plan_raw"], o["fact_raw"],
                o["plan"], o["fact"], f"=IF(AND(ISNUMBER(Q{i}),ISNUMBER(R{i})),1,0)", f'=IF(S{i}=1,R{i}-Q{i},"")',
                o["akc_status"], "", f"Скважины!{o['src_cells']}".replace("Скважины!", "реестр ")]
        for j, v in enumerate(vals, 1):
            c = wo.cell(row=i, column=j, value=v)
            c.font = F(size=9)
        checks.append(("Операции", f"S{i}", o["pair"]))
        if o["pair"]:
            checks.append(("Операции", f"T{i}", o["dev"]))
    for j, wdt in enumerate([12, 20, 6, 7, 9, 10, 8, 16, 8, 9, 10, 8, 10, 8, 10, 10, 9, 9, 8, 11, 22, 11, 16], 1):
        wo.column_dimensions[get_column_letter(j)].width = wdt
    wo.freeze_panes = "B2"
    L = N + 1
    R = lambda col: f"Операции!${col}$2:${col}${L}"

    # --- помощники формул ---
    def crit(cond):
        """cond: dict колонка→значение; всегда база=1 и цементируется=1"""
        c = {"G": 1, "J": 1, **cond}
        ifs = ",".join(f'{R(k)},{json.dumps(v, ensure_ascii=False) if isinstance(v, str) else v}' for k, v in c.items())
        arr = "*".join(f'({R(k)}={json.dumps(v, ensure_ascii=False) if isinstance(v, str) else v})' for k, v in c.items())
        return ifs, arr

    def stat_row(ws, row, col, cond, valcol, exp, kinds=("n", "mean", "median", "min", "max")):
        ifs, arr = crit(cond)
        out = {}
        for j, kind in enumerate(kinds):
            cell = ws.cell(row=row, column=col + j)
            ref = f"{get_column_letter(col + j)}{row}"
            if kind == "n":
                cell.value = f'=COUNTIFS({ifs},{R(valcol)},">-1E+99")'
            elif kind == "mean":
                cell.value = f'=IFERROR(AVERAGEIFS({R(valcol)},{ifs}),"н/д")'
            elif kind == "median":
                cell.value = ArrayFormula(ref, f'=IFERROR(MEDIAN(IF({arr}*ISNUMBER({R(valcol)}),{R(valcol)})),"н/д")')
            elif kind == "min":
                cell.value = f'=IF({get_column_letter(col)}{row}=0,"н/д",_xlfn.MINIFS({R(valcol)},{ifs}))'
            elif kind == "max":
                cell.value = f'=IF({get_column_letter(col)}{row}=0,"н/д",_xlfn.MAXIFS({R(valcol)},{ifs}))'
            cell.font = F(size=10)
            cell.number_format = "0" if kind == "n" else ("0.00" if valcol in ("L", "N") else "0.0")
            v = exp.get(kind)
            checks.append((ws.title, ref, v if v is not None else ("н/д" if kind != "n" else 0)))
        return out

    def label(ws, row, text, bold=False, fill=None, col=1):
        c = ws.cell(row=row, column=col, value=text)
        c.font = F(bold=bold, size=10, color="FFFFFF" if fill else "00326E")
        if fill:
            c.fill = fill
        return c

    def count_cell(ws, ref, formula, exp):
        ws[ref] = formula
        ws[ref].font = F(size=10)
        checks.append((ws.title, ref, exp))

    # ---- S1_Охват
    s = wb.create_sheet("S1_Охват")
    label(s, 1, "S1 · Охват анализа (слайд 2)", bold=True)
    cov = M["coverage"]
    items = [
        ("Записей (скважин) в реестре", '=COUNTA(Скважины!$A$2:$A$200)', cov["registry_wells"]),
        ("Скважин, законченных бурением в 2026 г. (база)", '=COUNTIFS(Скважины!$J$2:$J$200,2026)', cov["wells"]),
        ("Записей без даты окончания бурения", '=COUNTIFS(Скважины!$J$2:$J$200,"н/д")', cov["wells_nodate"]),
        ("База: наклонно направленных (НН)", '=COUNTIFS(Скважины!$J$2:$J$200,2026,Скважины!$D$2:$D$200,"НН")', cov["wells_nn"]),
        ("База: горизонтальных (ГС)", '=COUNTIFS(Скважины!$J$2:$J$200,2026,Скважины!$D$2:$D$200,"ГС")', cov["wells_gs"]),
        ("База: колонн всего (3 на скважину)", f'=COUNTIFS({R("G")},1)', cov["strings_total"]),
        ("База: нецементируемых колонн («н/ц»)", f'=COUNTIFS({R("G")},1,{R("J")},0)', cov["strings_uncemented"]),
        ("База: цементируемых колонн (операций)", f'=COUNTIFS({R("G")},1,{R("J")},1)', cov["strings_cemented"]),
        ("…с Кобщ ОС", f'=COUNTIFS({R("G")},1,{R("J")},1,{R("L")},">-1E+99")', cov["strings_akc_oh"]),
        ("…с Кобщ по всему интервалу", f'=COUNTIFS({R("G")},1,{R("J")},1,{R("N")},">-1E+99")', cov["strings_akc_full"]),
        ("…без АКЦ (н/д)", f'=B10-SUMPRODUCT(({R("G")}=1)*({R("J")}=1)*((ISNUMBER({R("L")})+ISNUMBER({R("N")}))>0))', cov["strings_akc_missing"]),
        ("…с сопоставимой парой ВПЦ план/факт", f'=SUMIFS({R("S")},{R("G")},1,{R("J")},1)', cov["toc_pairs"]),
    ]
    hdr(s, 2, ["Показатель", "Значение"])
    for i, (t, f_, e) in enumerate(items, 3):
        s.cell(row=i, column=1, value=t).font = F(size=10)
        count_cell(s, f"B{i}", f_, e)
    s.column_dimensions["A"].width = 52
    s.column_dimensions["B"].width = 12
    # полнота по колоннам
    r0 = len(items) + 5
    label(s, r0, "Полнота данных по колоннам (база 2026, цементируемые колонны)", bold=True)
    hdr(s, r0 + 1, ["Колонна", "Колонн", "Кобщ ОС", "Кобщ весь", "Пара ВПЦ", "Без АКЦ"])
    for i, (k, name, *_ ) in enumerate(STRINGS):
        rr = r0 + 2 + i
        s.cell(row=rr, column=1, value=name).font = F(size=10)
        sel = [o for o in ops if o["base"] and o["cemented"] and o["key"] == k]
        base = f'{R("G")},1,{R("J")},1,{R("H")},"{name}"'
        count_cell(s, f"B{rr}", f"=COUNTIFS({base})", len(sel))
        count_cell(s, f"C{rr}", f'=COUNTIFS({base},{R("L")},">-1E+99")', len([o for o in sel if o["os"] is not None]))
        count_cell(s, f"D{rr}", f'=COUNTIFS({base},{R("N")},">-1E+99")', len([o for o in sel if o["full"] is not None]))
        count_cell(s, f"E{rr}", f'=SUMIFS({R("S")},{base})', sum(o["pair"] for o in sel))
        count_cell(s, f"F{rr}", f'=B{rr}-SUMPRODUCT(({R("G")}=1)*({R("J")}=1)*({R("H")}="{name}")*((ISNUMBER({R("L")})+ISNUMBER({R("N")}))>0))',
                   len([o for o in sel if o["os"] is None and o["full"] is None]))

    # ---- S2_Итоги по колоннам
    s = wb.create_sheet("S2_Колонны")
    label(s, 1, "S2 · Итоги 2026 по колоннам: Кобщ ОС, Кобщ весь интервал, ВПЦ (слайд 3). 2025 г. — данных нет.", bold=True)
    row = 3
    for valcol, title in (("L", "Кобщ, открытый ствол"), ("N", "Кобщ, весь интервал"), ("T", "ВПЦ: отклонение факт − план, м (только пары)")):
        label(s, row, title, bold=True)
        hdr(s, row + 1, ["Колонна", "n", "Среднее", "Медиана", "Мин", "Макс"])
        for i, (k, name, *_ ) in enumerate(STRINGS):
            rr = row + 2 + i
            s.cell(row=rr, column=1, value=name).font = F(size=10)
            key = {"L": "bond_oh", "N": "bond_full", "T": "toc"}[valcol]
            stat_row(s, rr, 2, {"H": name}, valcol, M["by_string"][k][key])
        row += 6
    label(s, row, "Справочно: доля значений не ниже порогов (критерий не утверждён)", bold=True)
    hdr(s, row + 1, ["Колонна · показатель", "n", "≥ 0,80", "≥ 0,90"])
    rr = row + 2
    for k, name, *_ in STRINGS:
        for valcol, key, lab in (("L", "os", "ОС"), ("N", "full", "весь")):
            s.cell(row=rr, column=1, value=f"{name} · {lab}").font = F(size=10)
            base = f'{R("G")},1,{R("J")},1,{R("H")},"{name}"'
            t = M["thresholds"][f"{k}_{key}"]
            count_cell(s, f"B{rr}", f'=COUNTIFS({base},{R(valcol)},">-1E+99")', t["n"])
            count_cell(s, f"C{rr}", f'=COUNTIFS({base},{R(valcol)},">=0.8")', t["ge80"])
            count_cell(s, f"D{rr}", f'=COUNTIFS({base},{R(valcol)},">=0.9")', t["ge90"])
            rr += 1
    row = rr + 1
    label(s, row, "Динамика внутри 2026 г.: медиана Кобщ весь интервал по кварталу окончания бурения", bold=True)
    hdr(s, row + 1, ["Квартал · колонна", "n", "Среднее", "Медиана", "Мин", "Макс"])
    rr = row + 2
    for q in M["quarters"]:
        for k, name, *_ in STRINGS:
            s.cell(row=rr, column=1, value=f"{q} · {name}").font = F(size=10)
            stat_row(s, rr, 2, {"F": q, "H": name}, "N", M["by_quarter"][q][k]["bond_full"])
            rr += 1
    s.column_dimensions["A"].width = 34

    # ---- S3_Тип_Колонна
    s = wb.create_sheet("S3_Тип_Колонна")
    label(s, 1, "S3 · Тип скважины × колонна (слайд 4)", bold=True)
    hdr(s, 2, ["Профиль · колонна", "n ОС", "ОС ср.", "ОС мед.", "ОС мин", "ОС макс", "n весь", "Весь ср.", "Весь мед.", "Весь мин", "Весь макс",
               "Пар ВПЦ", "Откл. ср., м", "Откл. мед., м", "Откл. мин", "Откл. макс"])
    rr = 3
    for p in ("НН", "ГС"):
        for k, name, *_ in STRINGS:
            s.cell(row=rr, column=1, value=f"{p} · {name}").font = F(size=10)
            g_ = M["by_type"][p][k]
            stat_row(s, rr, 2, {"D": p, "H": name}, "L", g_["bond_oh"])
            stat_row(s, rr, 7, {"D": p, "H": name}, "N", g_["bond_full"])
            stat_row(s, rr, 12, {"D": p, "H": name}, "T", g_["toc"])
            rr += 1
    s.column_dimensions["A"].width = 28

    # ---- S4_Месторождения
    s = wb.create_sheet("S4_Месторождения")
    label(s, 1, "S4 · Месторождения и буровые подрядчики (слайд 5)", bold=True)
    hdr(s, 2, ["Месторождение · колонна", "n ОС", "ОС ср.", "ОС мед.", "ОС мин", "ОС макс", "n весь", "Весь ср.", "Весь мед.", "Весь мин", "Весь макс", "Пар ВПЦ"])
    rr = 3
    for f_ in M["fields"]:
        for k, name, *_ in STRINGS:
            s.cell(row=rr, column=1, value=f"{f_} · {name}").font = F(size=10)
            g_ = M["by_field"][f_][k]
            stat_row(s, rr, 2, {"B": f_, "H": name}, "L", g_["bond_oh"])
            stat_row(s, rr, 7, {"B": f_, "H": name}, "N", g_["bond_full"])
            count_cell(s, f"L{rr}", f'=SUMIFS({R("S")},{R("G")},1,{R("J")},1,{R("B")},"{f_}",{R("H")},"{name}")', g_["toc"]["n"])
            rr += 1
    rr += 1
    label(s, rr, "Полнота по месторождениям (база 2026)", bold=True)
    hdr(s, rr + 1, ["Месторождение", "Скважин", "НН", "ГС", "Цем. колонн", "С Кобщ весь", "Пар ВПЦ"])
    rr += 2
    for f_ in M["fields"]:
        bf = M["by_field"][f_]
        s.cell(row=rr, column=1, value=f_).font = F(size=10)
        count_cell(s, f"B{rr}", f'=COUNTIFS(Скважины!$B$2:$B$200,"{f_}",Скважины!$J$2:$J$200,2026)', bf["wells"])
        count_cell(s, f"C{rr}", f'=COUNTIFS(Скважины!$B$2:$B$200,"{f_}",Скважины!$J$2:$J$200,2026,Скважины!$D$2:$D$200,"НН")', bf["wells_nn"])
        count_cell(s, f"D{rr}", f'=COUNTIFS(Скважины!$B$2:$B$200,"{f_}",Скважины!$J$2:$J$200,2026,Скважины!$D$2:$D$200,"ГС")', bf["wells_gs"])
        count_cell(s, f"E{rr}", f'=COUNTIFS({R("G")},1,{R("J")},1,{R("B")},"{f_}")', bf["strings_cemented"])
        count_cell(s, f"F{rr}", f'=COUNTIFS({R("G")},1,{R("J")},1,{R("B")},"{f_}",{R("N")},">-1E+99")', bf["strings_akc_full"])
        count_cell(s, f"G{rr}", f'=SUMIFS({R("S")},{R("G")},1,{R("J")},1,{R("B")},"{f_}")', bf["toc_pairs"])
        rr += 1
    rr += 1
    label(s, rr, "Буровой подрядчик × колонна: Кобщ весь интервал (подрядчик по цементированию в реестре отсутствует)", bold=True)
    hdr(s, rr + 1, ["Подрядчик · колонна", "n", "Среднее", "Медиана", "Мин", "Макс"])
    rr += 2
    for c_ in M["contractors"]:
        for k, name, *_ in STRINGS:
            s.cell(row=rr, column=1, value=f"{c_} · {name}").font = F(size=10)
            stat_row(s, rr, 2, {"E": c_, "H": name}, "N", M["by_contractor"][c_][k]["bond_full"])
            rr += 1
    s.column_dimensions["A"].width = 36

    # ---- S5 / S6: месторождение × тип × колонна
    s5 = wb.create_sheet("S5_Сцепление_разрез")
    s6 = wb.create_sheet("S6_ВПЦ_разрез")
    label(s5, 1, "S5 · Месторождение × тип × колонна: Кобщ (слайд 6)", bold=True)
    label(s6, 1, "S6 · Месторождение × тип × колонна: ВПЦ (слайд 7). Отклонение = факт − план; > 0 — недоподъём.", bold=True)
    hdr(s5, 2, ["Месторождение · профиль · колонна", "n ОС", "ОС ср.", "ОС мед.", "ОС мин", "ОС макс", "n весь", "Весь ср.", "Весь мед.", "Весь мин", "Весь макс"])
    hdr(s6, 2, ["Месторождение · профиль · колонна", "Пар", "Откл. ср.", "Откл. мед.", "Откл. мин", "Откл. макс", "План мед., м", "Факт мед., м", "Недоподъём, пар"])
    r5 = r6 = 3
    for f_ in M["fields"]:
        for p in ("НН", "ГС"):
            for k, name, *_ in STRINGS:
                g_ = M["by_field_type"][f_][p][k]
                if g_["n_strings"] == 0:
                    continue
                cond = {"B": f_, "D": p, "H": name}
                s5.cell(row=r5, column=1, value=f"{f_} · {p} · {name}").font = F(size=10)
                stat_row(s5, r5, 2, cond, "L", g_["bond_oh"])
                stat_row(s5, r5, 7, cond, "N", g_["bond_full"])
                r5 += 1
                s6.cell(row=r6, column=1, value=f"{f_} · {p} · {name}").font = F(size=10)
                stat_row(s6, r6, 2, cond, "T", g_["toc"])
                ifs, arr = crit(cond)
                pr = f"{arr}*({R('S')}=1)"
                s6[f"G{r6}"] = ArrayFormula(f"G{r6}", f'=IFERROR(MEDIAN(IF({pr},{R("Q")})),"н/д")')
                s6[f"H{r6}"] = ArrayFormula(f"H{r6}", f'=IFERROR(MEDIAN(IF({pr},{R("R")})),"н/д")')
                count_cell(s6, f"I{r6}", f'=COUNTIFS({ifs},{R("T")},">0")', g_["toc"]["under_n"])
                checks.append(("S6_ВПЦ_разрез", f"G{r6}", g_["toc"]["plan_median"] if g_["toc"]["plan_median"] is not None else "н/д"))
                checks.append(("S6_ВПЦ_разрез", f"H{r6}", g_["toc"]["fact_median"] if g_["toc"]["fact_median"] is not None else "н/д"))
                r6 += 1
    s5.column_dimensions["A"].width = 44
    s6.column_dimensions["A"].width = 44

    # ---- S7_Худшие + чувствительность
    s = wb.create_sheet("S7_Худшие")
    label(s, 1, "S7 · Наихудшие операции (слайд 8): сортировка по Кобщ весь интервал ↑ и по недоподъёму ↓. Значения — ссылки на «Операции».", bold=True)
    hdr(s, 2, ["Скважина", "Колонна", "Строка «Операции»", "Кобщ ОС", "Кобщ весь", "Отклонение ВПЦ, м"])
    idx = {(o["well"], o["string"]): i for i, o in enumerate(ops, 2)}
    rr = 3
    for w in M["worst_bond"]:
        i = idx[(w["well"], w["string"])]
        s.cell(row=rr, column=1, value=w["well"]); s.cell(row=rr, column=2, value=w["string"]); s.cell(row=rr, column=3, value=i)
        s[f"D{rr}"] = f"=Операции!L{i}"; s[f"E{rr}"] = f"=Операции!N{i}"; s[f"F{rr}"] = f'=IF(Операции!S{i}=1,Операции!T{i},"н/д")'
        checks.append(("S7_Худшие", f"E{rr}", w["full"]))
        rr += 1
    rr += 1
    hdr(s, rr, ["Скважина", "Колонна", "Строка «Операции»", "ВПЦ план, м", "ВПЦ факт, м", "Недоподъём, м"])
    rr += 1
    for w in M["worst_toc"]:
        i = idx[(w["well"], w["string"])]
        s.cell(row=rr, column=1, value=w["well"]); s.cell(row=rr, column=2, value=w["string"]); s.cell(row=rr, column=3, value=i)
        s[f"D{rr}"] = f"=Операции!Q{i}"; s[f"E{rr}"] = f"=Операции!R{i}"; s[f"F{rr}"] = f"=Операции!T{i}"
        checks.append(("S7_Худшие", f"F{rr}", w["dev"]))
        rr += 1
    rr += 1
    label(s, rr, "Чувствительность: исключений по невалидности нет (оснований в реестре нет). Вариант «база + записи без даты окончания бурения» (формулы без условия «База 2026»).", bold=True)
    hdr(s, rr + 1, ["Колонна · показатель", "База: n", "База: ср.", "База: мед.", "+ без даты: n", "+ без даты: ср.", "+ без даты: мед."])
    rr += 2
    for k, name, *_ in STRINGS:
        for valcol, key, lab in (("L", "os", "ОС"), ("N", "full", "весь")):
            s.cell(row=rr, column=1, value=f"{name} · {lab}")
            stat_row(s, rr, 2, {"H": name}, valcol, M["sensitivity"]["base"][k][key], kinds=("n", "mean", "median"))
            # без условия базы
            allf = f'{R("J")},1,{R("H")},"{name}"'
            arr = f'({R("J")}=1)*({R("H")}="{name}")'
            ex = M["sensitivity"]["with_nodate"][k][key]
            count_cell(s, f"E{rr}", f'=COUNTIFS({allf},{R(valcol)},">-1E+99")', ex["n"])
            s[f"F{rr}"] = f'=IFERROR(AVERAGEIFS({R(valcol)},{allf}),"н/д")'
            checks.append(("S7_Худшие", f"F{rr}", ex["mean"]))
            s[f"G{rr}"] = ArrayFormula(f"G{rr}", f'=IFERROR(MEDIAN(IF({arr}*ISNUMBER({R(valcol)}),{R(valcol)})),"н/д")')
            checks.append(("S7_Худшие", f"G{rr}", ex["median"]))
            for c in "EFG":
                s[f"{c}{rr}"].number_format = "0" if c == "E" else "0.00"
            rr += 1
    s.column_dimensions["A"].width = 24

    # ---- S8_Полнота
    s = wb.create_sheet("S8_Полнота")
    label(s, 1, "S8 · Полнота данных по скважинам и колоннам (слайд 9): статус исходной ячейки", bold=True)
    hdr(s, 2, ["Скважина", "Месторождение", "Профиль", "База 2026"] + [f"{n}: {x}" for _, n, _, _ in STRINGS for x in ("ОС", "весь", "ВПЦ пара")])
    rr = 3
    for w in wells:
        s.cell(row=rr, column=1, value=w["well"]); s.cell(row=rr, column=2, value=w["field"]); s.cell(row=rr, column=3, value=w["profile"]); s.cell(row=rr, column=4, value=w["base"])
        j = 5
        for k, *_ in STRINGS:
            o = next(o for o in ops if o["well"] == w["well"] and o["key"] == k and o["row"] == w["row"])
            for v in (o["os_raw"], o["full_raw"]):
                s.cell(row=rr, column=j, value=v if v is not None else "пусто"); j += 1
            s.cell(row=rr, column=j, value=("н/ц" if not o["cemented"] else ("да" if o["pair"] else "нет"))); j += 1
        rr += 1
    rr += 1
    label(s, rr, "Нет в реестре (н/д для всех записей):", bold=True)
    for t in M["missing_fields"]:
        rr += 1
        s.cell(row=rr, column=1, value="• " + t)
    s.column_dimensions["A"].width = 14
    s.column_dimensions["B"].width = 22

    # ---- Журнал_проверок
    s = wb.create_sheet("Журнал_проверок")
    hdr(s, 1, ["Тип", "Скважина", "Описание", "Ячейки реестра"])
    for i, row_ in enumerate(log, 2):
        for j, v in enumerate(row_, 1):
            c = s.cell(row=i, column=j, value=v)
            c.font = F(size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for j, wdt in enumerate([24, 22, 100, 16], 1):
        s.column_dimensions[get_column_letter(j)].width = wdt

    for w_ in wb.worksheets:
        w_.sheet_view.showGridLines = w_.title in ("Реестр_исходный", "Операции", "Скважины")
    wb.save(OUT_XLSX)
    return checks


def main():
    wells, ops, raw_rows, raw_wb = read_registry()
    M = compute(wells, ops)
    log = audit(wells, ops, raw_wb)
    M["log"] = log
    os.makedirs(os.path.dirname(OUT_XLSX), exist_ok=True)
    checks = write_book(wells, ops, M, log, raw_wb, raw_rows)
    json.dump(M, open(OUT_JSON, "w", encoding="utf-8"), ensure_ascii=False, indent=1, default=str)
    json.dump(checks, open(os.path.join(ROOT, "data", "checks.json"), "w", encoding="utf-8"), ensure_ascii=False, default=str)
    print(json.dumps(M["coverage"], ensure_ascii=False, indent=1))
    print("checks:", len(checks))


if __name__ == "__main__":
    main()
