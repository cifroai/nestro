# -*- coding: utf-8 -*-
"""Минимальный вычислитель формул Excel для проверки сгенерированной книги."""
import datetime as dt, re, sys
import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

EPOCH = dt.date(1899, 12, 30)

def to_serial(v):
    if isinstance(v, dt.datetime): v = v.date()
    if isinstance(v, dt.date): return (v - EPOCH).days
    return v

TOK = re.compile(r"""
    (?P<str>"(?:[^"]|"")*")
  | (?P<num>\d+(?:\.\d+)?)
  | (?P<ref>(?:'[^']+'|[A-Za-zА-Яа-яЁё_][A-Za-zА-Яа-яЁё0-9_.]*)!\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)
  | (?P<lref>\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)
  | (?P<fn>[A-Za-zА-Яа-яЁё_][A-Za-zА-Яа-яЁё0-9_.]*)\s*(?=\()
  | (?P<op><=|>=|<>|[-+*/<>=,()])
  | (?P<ws>\s+)
""", re.X)

def tokenize(s):
    out, i = [], 0
    while i < len(s):
        m = TOK.match(s, i)
        if not m: raise ValueError(f"tokenize fail at {i}: {s[i:i+25]!r} in {s!r}")
        k = m.lastgroup
        if k != "ws": out.append((k, m.group()))
        i = m.end()
    return out

class Ev:
    def __init__(self, wb):
        self.wb = wb
        self.cache = {}
        self.stack = []

    def cell(self, sheet, coord):
        coord = coord.replace("$", "")
        key = (sheet, coord)
        if key in self.cache: return self.cache[key]
        if key in self.stack: raise ValueError(f"circular {key}")
        self.stack.append(key)
        try:
            c = self.wb[sheet][coord]
            v = c.value
            if isinstance(v, str) and v.startswith("="):
                v = self.eval(v[1:], sheet)
            else:
                v = to_serial(v)
                if v is None: v = ""
        finally:
            self.stack.pop()
        self.cache[key] = v
        return v

    def rng(self, sheet, ref):
        ref = ref.replace("$", "")
        if ":" not in ref: return [self.cell(sheet, ref)]
        a, b = ref.split(":")
        c1, r1 = re.match(r"([A-Z]+)(\d+)", a).groups()
        c2, r2 = re.match(r"([A-Z]+)(\d+)", b).groups()
        out = []
        for rr in range(int(r1), int(r2) + 1):
            for cc in range(column_index_from_string(c1), column_index_from_string(c2) + 1):
                out.append(self.cell(sheet, f"{get_column_letter(cc)}{rr}"))
        return out

    def eval(self, expr, sheet):
        # состояние парсера сохраняем: вычисление ссылки рекурсивно вызывает eval
        saved = (getattr(self, "toks", None), getattr(self, "pos", None), getattr(self, "sheet", None))
        try:
            self.toks = tokenize(expr); self.pos = 0; self.sheet = sheet
            v = self.p_cmp()
            if self.pos != len(self.toks): raise ValueError(f"trailing tokens in {expr!r}")
            return v
        finally:
            self.toks, self.pos, self.sheet = saved

    def peek(self):
        return self.toks[self.pos] if self.pos < len(self.toks) else (None, None)

    def p_cmp(self):
        l = self.p_add()
        while self.peek()[1] in ("<", ">", "<=", ">=", "=", "<>"):
            op = self.toks[self.pos][1]; self.pos += 1
            r = self.p_add()
            if isinstance(l, str) or isinstance(r, str):
                ls, rs = str(l), str(r)
                l = {"=": ls == rs, "<>": ls != rs, "<": ls < rs, ">": ls > rs,
                     "<=": ls <= rs, ">=": ls >= rs}[op]
            else:
                l = {"=": l == r, "<>": l != r, "<": l < r, ">": l > r,
                     "<=": l <= r, ">=": l >= r}[op]
        return l

    def p_add(self):
        l = self.p_mul()
        while self.peek()[1] in ("+", "-"):
            op = self.toks[self.pos][1]; self.pos += 1
            r = self.p_mul()
            l = (l or 0) + (r or 0) if op == "+" else (l or 0) - (r or 0)
        return l

    def p_mul(self):
        l = self.p_un()
        while self.peek()[1] in ("*", "/"):
            op = self.toks[self.pos][1]; self.pos += 1
            r = self.p_un()
            l = l * r if op == "*" else l / r
        return l

    def p_un(self):
        if self.peek()[1] == "-":
            self.pos += 1; return -self.p_un()
        return self.p_atom()

    def p_atom(self):
        k, t = self.peek()
        if k == "num": self.pos += 1; return float(t) if "." in t else int(t)
        if k == "str": self.pos += 1; return t[1:-1].replace('""', '"')
        if k == "fn":
            self.pos += 1; name = t.upper()
            assert self.toks[self.pos][1] == "("; self.pos += 1
            args = []
            if self.peek()[1] != ")":
                while True:
                    sv = self.pos
                    kk, tt = self.peek()
                    if kk in ("ref", "lref") and ":" in tt and self.toks[self.pos+1][1] in (",", ")"):
                        self.pos += 1
                        if kk == "ref":
                            sh, rf = tt.split("!"); sh = sh.strip("'")
                        else:
                            sh, rf = self.sheet, tt
                        args.append(("RANGE", self.rng(sh, rf)))
                    else:
                        self.pos = sv
                        args.append(("V", self.p_cmp()))
                    if self.peek()[1] == ",": self.pos += 1; continue
                    break
            assert self.toks[self.pos][1] == ")", f"expected ) in {name}"
            self.pos += 1
            flat = []
            for kind, a in args:
                flat.extend(a) if kind == "RANGE" else flat.append(a)
            nums = [x for x in flat if isinstance(x, (int, float)) and not isinstance(x, bool)]
            if name == "SUM": return sum(nums)
            if name == "MIN": return min(nums) if nums else 0
            if name == "MAX": return max(nums) if nums else 0
            if name == "IF":
                vals = [a for _, a in args]
                return vals[1] if vals[0] else (vals[2] if len(vals) > 2 else False)
            if name == "TODAY": return to_serial(TODAY)
            if name == "COUNTIF":
                rngv, crit = args[0][1], args[1][1]
                return sum(1 for x in rngv if str(x) == str(crit))
            raise ValueError(f"unsupported fn {name}")
        if k == "ref":
            self.pos += 1
            sh, rf = t.split("!"); sh = sh.strip("'")
            return self.rng(sh, rf)[0] if ":" not in rf else self.rng(sh, rf)
        if k == "lref":
            self.pos += 1
            return self.cell(self.sheet, t) if ":" not in t else self.rng(self.sheet, t)
        if t == "(":
            self.pos += 1; v = self.p_cmp()
            assert self.toks[self.pos][1] == ")"; self.pos += 1; return v
        raise ValueError(f"unexpected token {t!r}")
