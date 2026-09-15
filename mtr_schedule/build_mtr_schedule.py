# -*- coding: utf-8 -*-
"""
График закупки и доставки МТР для крепления скважин ЮГ-1, ЮГ-2, ЮГ-3.

Логистическая схема унаследована от файла "График доставки цемента на скв. ГЕО-12":
завод-изготовитель (Тюмень/Урал) -> ж/д до Владивостока -> морской переход
до Петропавловска-Камчатского -> доставка с базы Заказчика до скважины.

Формирует книгу Excel с листами:
  1. Инструкция
  2. Справочник      - длительности этапов, программа крепления, номенклатура
  3..5. ЮГ-1/ЮГ-2/ЮГ-3 - подробный график по этапам для каждой номенклатуры
  6. СВОДНЫЙ ГРАФИК  - диаграмма Ганта с вертикальной линией текущей даты
  7. Сводная потребность
"""
import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.formatting.rule import Rule
from openpyxl.worksheet.datavalidation import DataValidation

# ----------------------------------------------------------------------------
# ИСХОДНЫЕ ДАННЫЕ (все допущения вынесены сюда и продублированы на листе
# "Инструкция" — меняются в одном месте)
# ----------------------------------------------------------------------------

TODAY = dt.date(2026, 9, 15)          # дата формирования графика
PROCUREMENT_READY = dt.date(2026, 9, 20)  # ранее этой даты заказ разместить невозможно
TARGET_BUFFER = 14                    # плановый резерв времени к требуемому сроку, сут
RISK_THRESHOLD = 14                   # запас < 14 сут -> РИСК
NEED_LEAD = 5                         # МТР на скважине за 5 сут до операции

DATE_FMT = "DD.MM.YYYY"

# --- Этапы поставки -----------------------------------------------------------
STAGES = [
    "Размещение заказа, заключение договора, согласование спецификации",
    "Изготовление на заводе-изготовителе (для тампонажных материалов — приготовление и замес смеси)",
    "Подготовка к отгрузке: растаривание/затарка в МКР, пакетирование, маркировка",
    "Перевозка на контейнерный терминал, затарка в контейнеры, погрузка на ж/д",
    "Перевозка ж/д транспортом (завод-изготовитель — ст. Владивосток)",
    "Перевалка во Владивостоке на морское судно",
    "Морской переход Владивосток — Петропавловск-Камчатский",
    "Доставка с базы Заказчика в г. Петропавловск-Камчатский до скважины",
]

# --- Категории МТР и длительности этапов, суток -------------------------------
CATEGORIES = [
    ("CEM",  "Тампонажные материалы (цементы, смеси)", [10, 3, 2, 2, 14, 2, 7, 4]),
    ("CHEM", "Химреагенты (замедлитель)",              [10, 5, 1, 1, 14, 2, 7, 4]),
    ("OK",   "Трубы обсадные",                         [14, 35, 2, 3, 18, 3, 7, 5]),
    ("TO",   "Технологическая оснастка",               [10, 25, 2, 2, 16, 2, 7, 4]),
    ("SPEC", "Спецоборудование (пакер, переводник)",   [14, 60, 2, 2, 16, 2, 7, 4]),
]
CAT_COL = {code: 3 + i for i, (code, _, _) in enumerate(CATEGORIES)}  # C..G
CAT_NAME = {code: name for code, name, _ in CATEGORIES}
CAT_DUR = {code: dur for code, _, dur in CATEGORIES}

# --- Программа крепления скважин ---------------------------------------------
# (скважина, этап крепления, дата начала операции по ГТН)
OPERATIONS = [
    ("ЮГ-1", "Направление ОК-426",        dt.date(2026, 11, 20)),
    ("ЮГ-1", "Кондуктор ОК-324",          dt.date(2026, 12,  5)),
    ("ЮГ-1", "Экспл. колонна ОК-245",     dt.date(2027,  1,  5)),
    ("ЮГ-1", "Хвостовик ОК-168",          dt.date(2027,  1, 20)),
    ("ЮГ-2", "Направление ОК-426",        dt.date(2027,  1, 19)),
    ("ЮГ-2", "Кондуктор ОК-324",          dt.date(2027,  2,  3)),
    ("ЮГ-2", "Экспл. колонна ОК-245",     dt.date(2027,  3,  6)),
    ("ЮГ-2", "Хвостовик ОК-168",          dt.date(2027,  3, 21)),
    ("ЮГ-3", "Направление ОК-426",        dt.date(2027,  3, 20)),
    ("ЮГ-3", "Кондуктор ОК-324",          dt.date(2027,  4,  4)),
    ("ЮГ-3", "Экспл. колонна ОК-245",     dt.date(2027,  5,  5)),
    ("ЮГ-3", "Хвостовик ОК-168",          dt.date(2027,  5, 20)),
]
WELLS = ["ЮГ-1", "ЮГ-2", "ЮГ-3"]
OP_ROW = {}   # (well, op) -> строка на листе "Справочник"

# --- Номенклатура -------------------------------------------------------------
# (наименование, ед.изм, категория, этап крепления, кол-во на скважину, примечание)
NOMENCLATURE = [
    ("Цемент марки G (направление)", "т", "CEM", "Направление ОК-426", 15,
     "ПЦТ G HSR. Расчет: кольцевое пространство 590/426 мм, 0-50 м + стакан, к-т каверн. 1,3, запас 10%"),
    ("Цементная смесь тяжелая (кондуктор)", "т", "CEM", "Кондуктор ОК-324", 50,
     "ρ=1,85 г/см3. Кольцевое пространство 393,7/324 мм, 0-500 м, к-т каверн. 1,3, запас 10%"),
    ("Цементная смесь тяжелая (экспл.колонна)", "т", "CEM", "Экспл. колонна ОК-245", 32,
     "ρ=1,85 г/см3. Нижняя ступень 1200-1800 м, кольц. пр-во 295,3/245 мм, к-т каверн. 1,25"),
    ("Цементная смесь легкая (экспл.колонна)", "т", "CEM", "Экспл. колонна ОК-245", 45,
     "ρ=1,50 г/см3. Верхняя ступень 0-1200 м, кольц. пр-во 295,3/245 мм, к-т каверн. 1,25"),
    ("Замедлитель", "кг", "CHEM", "Направление ОК-426", 600,
     "Термостойкий замедлитель схватывания, 0,4% от массы цемента (142 т). Поставка единой партией под все этапы крепления"),
    ("Труба обсадная ОК-426", "м", "OK", "Направление ОК-426", 60,
     "426х10 мм, гр.пр. Д, резьба ОТТМ. Глубина спуска 50 м + технологический запас"),
    ("Труба обсадная ОК-324", "м", "OK", "Кондуктор ОК-324", 520,
     "324х9,5 мм, гр.пр. Д, резьба ОТТМ. Глубина спуска 500 м + технологический запас"),
    ("Труба обсадная ОК-245", "м", "OK", "Экспл. колонна ОК-245", 1830,
     "245х10 мм, гр.пр. Е, резьба ОТТГ. Глубина спуска 1800 м + технологический запас"),
    ("Тех.оснастка (ОК-426)", "компл.", "TO", "Направление ОК-426", 1,
     "Башмак БКМ-426, ЦКОД-426, центраторы ЦЦ-426/590 - 5 шт, пробка продавочная ПП-426"),
    ("Тех.оснастка (ОК-324)", "компл.", "TO", "Кондуктор ОК-324", 1,
     "Башмак БКМ-324, ЦКОД-324, центраторы ЦЦ-324/394 - 25 шт, ЦТ, пробка продавочная ПП-324"),
    ("Тех.оснастка (ОК-245)", "компл.", "TO", "Экспл. колонна ОК-245", 1,
     "Башмак БКМ-245, ЦКОД-245, МСЦ-245, центраторы ЦЦ-245/295 - 60 шт, пробки продавочные"),
    ("Тех.оснастка (ОК-168)", "компл.", "TO", "Хвостовик ОК-168", 1,
     "Башмак БКМ-168, ЦКОДМ-168, центраторы ЦЦ-168/216 - 30 шт, подвеска хвостовика"),
    ("Пакер набухающий термостойкий", "шт", "SPEC", "Хвостовик ОК-168", 2,
     "Набухающий в углеводородной/водной среде, рабочая температура до 200 С. Изоляция интервала хвостовика"),
    ("Разъединительный переводник 245/168", "шт", "SPEC", "Хвостовик ОК-168", 1,
     "Разъединитель подвески хвостовика 245/168, термостойкое исполнение. Длит. цикл изготовления - спецзаказ"),
]

# ----------------------------------------------------------------------------
# СТИЛИ
# ----------------------------------------------------------------------------
C_HEADER = "1F4E79"
C_SUBHEAD = "2E75B6"
C_ITEM = "D9E2F3"
C_INPUT = "FFF2CC"
C_BAND = "F7F9FC"

F_TITLE = Font(name="Calibri", size=14, bold=True, color="1F4E79")
F_SUB = Font(name="Calibri", size=10, italic=True, color="595959")
F_HDR = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
F_ITEM = Font(name="Calibri", size=10, bold=True, color="1F4E79")
F_BASE = Font(name="Calibri", size=9)
F_BASE_B = Font(name="Calibri", size=9, bold=True)
F_INPUT = Font(name="Calibri", size=9, bold=True, color="0000C0")

FILL_HDR = PatternFill("solid", fgColor=C_HEADER)
FILL_SUB = PatternFill("solid", fgColor=C_SUBHEAD)
FILL_ITEM = PatternFill("solid", fgColor=C_ITEM)
FILL_INPUT = PatternFill("solid", fgColor=C_INPUT)
FILL_BAND = PatternFill("solid", fgColor=C_BAND)

_thin = Side(style="thin", color="B4C6E7")
_med = Side(style="medium", color="1F4E79")
B_ALL = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
B_ITEM = Border(left=_med, right=_med, top=_med, bottom=_thin)

A_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
A_L = Alignment(horizontal="left", vertical="center", wrap_text=True)
A_LT = Alignment(horizontal="left", vertical="top", wrap_text=True)

# статусные цвета
CLR = {
    "В ГРАФИКЕ":     {"line": "00B050", "transport": "A9D08E", "prod": "E2EFDA", "text": "006100", "cell": "C6EFCE"},
    "РИСК":          {"line": "FFC000", "transport": "FFD966", "prod": "FFF2CC", "text": "9C5700", "cell": "FFEB9C"},
    "СРЫВ ГРАФИКА":  {"line": "FF0000", "transport": "F4B183", "prod": "FCE4D6", "text": "9C0006", "cell": "FFC7CE"},
}

wb = Workbook()

def style_range(ws, rng, font=None, fill=None, border=None, align=None, numfmt=None):
    for row in ws[rng]:
        for c in row:
            if font: c.font = font
            if fill: c.fill = fill
            if border: c.border = border
            if align: c.alignment = align
            if numfmt: c.number_format = numfmt


# ============================================================================
# ЛИСТ 1 — ИНСТРУКЦИЯ
# ============================================================================
ws_info = wb.active
ws_info.title = "Инструкция"
ws_info.sheet_view.showGridLines = False
ws_info.column_dimensions["A"].width = 3
ws_info.column_dimensions["B"].width = 42
ws_info.column_dimensions["C"].width = 96

r = 2
ws_info.cell(r, 2, "ГРАФИК ЗАКУПКИ И ДОСТАВКИ МТР ДЛЯ КРЕПЛЕНИЯ СКВАЖИН ЮГ-1, ЮГ-2, ЮГ-3").font = F_TITLE
ws_info.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
r += 1
ws_info.cell(r, 2, f"Дата формирования: {TODAY.strftime('%d.%m.%Y')}. "
                   "Логистическая схема унаследована из файла «График доставки цемента на скв. ГЕО-12».").font = F_SUB
ws_info.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)

def info_head(title):
    global r
    r += 2
    c = ws_info.cell(r, 2, title)
    c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    c.fill = FILL_SUB
    c.alignment = A_L
    ws_info.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws_info.cell(r, 3).fill = FILL_SUB

def info_row(k, v):
    global r
    r += 1
    a = ws_info.cell(r, 2, k); a.font = F_BASE_B; a.alignment = A_LT; a.border = B_ALL
    b = ws_info.cell(r, 3, v); b.font = F_BASE; b.alignment = A_LT; b.border = B_ALL

info_head("СОСТАВ КНИГИ")
info_row("Справочник",
         "Единая база модели: длительности 8 этапов поставки по категориям МТР, программа крепления скважин "
         "и номенклатура с объемами. Все листы считаются формулами от этого листа — правка здесь пересчитывает всю книгу.")
info_row("ЮГ-1 / ЮГ-2 / ЮГ-3",
         "Подробный график по каждой скважине: по каждой номенклатуре — ед.изм., кол-во, дата потребности "
         "и сроки (начало/окончание) по каждому из 8 этапов закупки и доставки.")
info_row("СВОДНЫЙ ГРАФИК",
         "Диаграмма Ганта по всем 42 позициям (14 номенклатур х 3 скважины) с вертикальной линией по текущей дате. "
         "Линия окрашивается по статусу строки: зеленая — в графике, желтая — риск, красная — срыв графика.")
info_row("Сводная потребность",
         "Суммарный объем закупки по номенклатуре на три скважины, самая ранняя дата потребности и худший статус.")

info_head("ЛОГИКА РАСЧЕТА")
info_row("Дата потребности",
         f"Дата операции по креплению минус {NEED_LEAD} сут (МТР должен быть на скважине заблаговременно). "
         "Задается на листе «Справочник», блок 2.")
info_row("Требуемая дата старта",
         "Дата потребности минус суммарная длительность цикла закупки и доставки. Крайний срок размещения заказа.")
info_row("Дата начала (план/факт)",
         f"Плановая дата размещения заказа. Принята как «требуемая дата старта минус {TARGET_BUFFER} сут резерва», "
         f"но не ранее {PROCUREMENT_READY.strftime('%d.%m.%Y')} (готовность закупочной процедуры). "
         "ЖЕЛТЫЕ ЯЧЕЙКИ — поле ввода: подставьте фактическую дату заключения договора, книга пересчитается.")
info_row("Прогноз поставки на скважину",
         "Дата начала (план/факт) + длительности всех 8 этапов последовательно.")
info_row("Запас, сут",
         "Дата потребности минус прогноз поставки. Отрицательное значение — опоздание.")

info_head("ЛЕГЕНДА СТАТУСОВ")
for st, txt in [("В ГРАФИКЕ", f"Запас >= {RISK_THRESHOLD} сут. Поставка обеспечивает срок крепления."),
                ("РИСК",      f"Запас от 0 до {RISK_THRESHOLD - 1} сут. Резерв времени исчерпан, любая задержка ведет к срыву."),
                ("СРЫВ ГРАФИКА", "Запас < 0. Прогнозная поставка позже даты потребности — требуются компенсирующие мероприятия.")]:
    r += 1
    a = ws_info.cell(r, 2, st)
    a.font = Font(name="Calibri", size=9, bold=True, color=CLR[st]["text"])
    a.fill = PatternFill("solid", fgColor=CLR[st]["cell"])
    a.alignment = A_C; a.border = B_ALL
    b = ws_info.cell(r, 3, txt); b.font = F_BASE; b.alignment = A_LT; b.border = B_ALL

info_head("ПРИНЯТЫЕ ДОПУЩЕНИЯ (подлежат уточнению Заказчиком)")
for k, v in [
    ("Конструкция скважин",
     "Направление ОК-426 (0-50 м), кондуктор ОК-324 (0-500 м), эксплуатационная колонна ОК-245 (0-1800 м), "
     "хвостовик ОК-168 (1700-2500 м). Конструкция принята одинаковой для ЮГ-1, ЮГ-2, ЮГ-3."),
    ("Программа бурения",
     "Скважины строятся последовательно одной буровой установкой со сдвигом 60 суток между скважинами. "
     "Даты операций по креплению — блок 2 листа «Справочник»."),
    ("Логистическая схема",
     "Завод-изготовитель (Тюмень/Урал) — ж/д до ст. Владивосток — перевалка на судно — морской переход "
     "до Петропавловска-Камчатского — доставка автотранспортом с базы Заказчика до куста. "
     "Длительности ж/д и морского плеча приняты по фактическим данным графика ГЕО-12."),
    ("Объемы тампонажных материалов",
     "Расчет по кольцевому пространству с коэффициентом кавернозности 1,25-1,30 и запасом 10%. "
     "Подлежат уточнению после кавернометрии и утверждения плана работ по креплению."),
    ("Труба обсадная ОК-168",
     "В перечне заявлена тех.оснастка ОК-168 и разъединительный переводник 245/168, но сама труба ОК-168 "
     "не заявлена. Принято, что хвостовик 168 мм обеспечивается из давальческих запасов Заказчика. "
     "Требует подтверждения — при закупке добавить позицию в номенклатуру (блок 3 листа «Справочник»)."),
    ("Замедлитель",
     "Закупается единой партией на все этапы крепления скважины, дата потребности привязана к самой ранней "
     "операции (крепление направления)."),
]:
    info_row(k, v)

info_head("КРИТИЧЕСКИЕ ПОЗИЦИИ")
info_row("Что требует решения сейчас",
         "По скважине ЮГ-1 позиции с длинным циклом изготовления (трубы ОК-426 и ОК-324, тех.оснастка ОК-426) "
         "имеют отрицательный запас: крайний срок размещения заказа уже прошел относительно даты формирования графика. "
         "Компенсирующие мероприятия указаны в столбце «Примечание / мероприятия» на листе скважины и в сводном графике.")

ws_info.freeze_panes = "A5"


# ============================================================================
# ЛИСТ 2 — СПРАВОЧНИК
# ============================================================================
ws_ref = wb.create_sheet("Справочник")
ws_ref.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGHIJ", [6, 62, 15, 15, 15, 15, 15, 14, 14, 60]):
    ws_ref.column_dimensions[col].width = w

ws_ref.cell(1, 1, "СПРАВОЧНИК ИСХОДНЫХ ДАННЫХ").font = F_TITLE
ws_ref.cell(2, 1, "Единственное место ввода. Все графики по скважинам и сводный график пересчитываются формулами от этого листа.").font = F_SUB

# --- Блок 1: длительности этапов
R1 = 4
c = ws_ref.cell(R1, 1, "1. ДЛИТЕЛЬНОСТЬ ЭТАПОВ ЗАКУПКИ И ДОСТАВКИ ПО КАТЕГОРИЯМ МТР, суток")
c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ws_ref.merge_cells(start_row=R1, start_column=1, end_row=R1, end_column=10)
for cc in range(1, 11):
    ws_ref.cell(R1, cc).fill = FILL_SUB
c.alignment = A_L

HDR1 = R1 + 1
ws_ref.cell(HDR1, 1, "№")
ws_ref.cell(HDR1, 2, "Наименование этапа")
for code, name, _ in CATEGORIES:
    ws_ref.cell(HDR1, CAT_COL[code], name)
ws_ref.row_dimensions[HDR1].height = 58
style_range(ws_ref, f"A{HDR1}:G{HDR1}", font=F_HDR, fill=FILL_HDR, border=B_ALL, align=A_C)

STAGE_R0 = HDR1 + 1   # первая строка этапа (этап 1)
for i, st in enumerate(STAGES):
    rr = STAGE_R0 + i
    ws_ref.cell(rr, 1, i + 1)
    ws_ref.cell(rr, 2, st)
    for code, _, dur in CATEGORIES:
        ws_ref.cell(rr, CAT_COL[code], dur[i])
    ws_ref.row_dimensions[rr].height = 26
    style_range(ws_ref, f"A{rr}:G{rr}", font=F_BASE, border=B_ALL, align=A_C)
    ws_ref.cell(rr, 2).alignment = A_L
    for code, _, _ in CATEGORIES:
        cell = ws_ref.cell(rr, CAT_COL[code])
        cell.fill = FILL_INPUT
        cell.font = F_INPUT

TOT_R = STAGE_R0 + len(STAGES)
ws_ref.cell(TOT_R, 2, "ИТОГО длительность цикла, суток")
for code, _, _ in CATEGORIES:
    cl = get_column_letter(CAT_COL[code])
    ws_ref.cell(TOT_R, CAT_COL[code], f"=SUM({cl}{STAGE_R0}:{cl}{TOT_R-1})")
style_range(ws_ref, f"A{TOT_R}:G{TOT_R}", font=F_BASE_B, fill=FILL_ITEM, border=B_ALL, align=A_C)
ws_ref.cell(TOT_R, 2).alignment = A_L

def dur_ref(code, stage_idx):
    """Абсолютная ссылка на длительность этапа (0-based) для категории."""
    return f"Справочник!${get_column_letter(CAT_COL[code])}${STAGE_R0 + stage_idx}"

# --- Блок 2: программа крепления
R2 = TOT_R + 2
c = ws_ref.cell(R2, 1, "2. ПРОГРАММА КРЕПЛЕНИЯ СКВАЖИН И ДАТЫ ПОТРЕБНОСТИ МТР")
c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ws_ref.merge_cells(start_row=R2, start_column=1, end_row=R2, end_column=10)
for cc in range(1, 11):
    ws_ref.cell(R2, cc).fill = FILL_SUB
c.alignment = A_L

HDR2 = R2 + 1
for col, txt in [(1, "№"), (2, "Скважина / этап крепления"), (3, "Дата операции по ГТН"),
                 (4, "Заблаговр., сут"), (5, "ДАТА ПОТРЕБНОСТИ (МТР на скважине)")]:
    ws_ref.cell(HDR2, col, txt)
ws_ref.row_dimensions[HDR2].height = 44
style_range(ws_ref, f"A{HDR2}:E{HDR2}", font=F_HDR, fill=FILL_HDR, border=B_ALL, align=A_C)

OP_R0 = HDR2 + 1
for i, (well, op, date) in enumerate(OPERATIONS):
    rr = OP_R0 + i
    OP_ROW[(well, op)] = rr
    ws_ref.cell(rr, 1, i + 1)
    ws_ref.cell(rr, 2, f"{well} — {op}")
    ws_ref.cell(rr, 3, date)
    ws_ref.cell(rr, 4, NEED_LEAD)
    ws_ref.cell(rr, 5, f"=C{rr}-D{rr}")
    style_range(ws_ref, f"A{rr}:E{rr}", font=F_BASE, border=B_ALL, align=A_C)
    ws_ref.cell(rr, 2).alignment = A_L
    ws_ref.cell(rr, 2).font = F_BASE_B
    for cc in (3, 4):
        ws_ref.cell(rr, cc).fill = FILL_INPUT
        ws_ref.cell(rr, cc).font = F_INPUT
    ws_ref.cell(rr, 3).number_format = DATE_FMT
    ws_ref.cell(rr, 5).number_format = DATE_FMT
    ws_ref.cell(rr, 5).font = F_BASE_B
    ws_ref.cell(rr, 5).fill = FILL_ITEM
    if i % 4 == 3:
        for cc in range(1, 6):
            ws_ref.cell(rr, cc).border = Border(left=_thin, right=_thin, top=_thin,
                                                bottom=Side(style="medium", color="1F4E79"))

def need_ref(well, op):
    return f"Справочник!$E${OP_ROW[(well, op)]}"

# --- Блок 3: номенклатура
R3 = OP_R0 + len(OPERATIONS) + 1
c = ws_ref.cell(R3, 1, "3. НОМЕНКЛАТУРА И ОБЪЕМЫ ЗАКУПКИ")
c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ws_ref.merge_cells(start_row=R3, start_column=1, end_row=R3, end_column=10)
for cc in range(1, 11):
    ws_ref.cell(R3, cc).fill = FILL_SUB
c.alignment = A_L

HDR3 = R3 + 1
for col, txt in [(1, "№"), (2, "Наименование МТР"), (3, "Ед. изм."), (4, "Категория поставки"),
                 (5, "Этап крепления"), (6, "ЮГ-1"), (7, "ЮГ-2"), (8, "ЮГ-3"), (9, "ИТОГО"),
                 (10, "Обоснование объема / характеристика")]:
    ws_ref.cell(HDR3, col, txt)
ws_ref.row_dimensions[HDR3].height = 34
style_range(ws_ref, f"A{HDR3}:J{HDR3}", font=F_HDR, fill=FILL_HDR, border=B_ALL, align=A_C)

NOM_R0 = HDR3 + 1
for i, (name, uom, cat, op, qty, note) in enumerate(NOMENCLATURE):
    rr = NOM_R0 + i
    ws_ref.cell(rr, 1, i + 1)
    ws_ref.cell(rr, 2, name)
    ws_ref.cell(rr, 3, uom)
    ws_ref.cell(rr, 4, CAT_NAME[cat])
    ws_ref.cell(rr, 5, op)
    for j in range(3):
        ws_ref.cell(rr, 6 + j, qty)
    ws_ref.cell(rr, 9, f"=SUM(F{rr}:H{rr})")
    ws_ref.cell(rr, 10, note)
    ws_ref.row_dimensions[rr].height = 32
    style_range(ws_ref, f"A{rr}:J{rr}", font=F_BASE, border=B_ALL, align=A_C)
    for cc in (2, 4, 5, 10):
        ws_ref.cell(rr, cc).alignment = A_L
    ws_ref.cell(rr, 2).font = F_BASE_B
    for j in range(3):
        ws_ref.cell(rr, 6 + j).fill = FILL_INPUT
        ws_ref.cell(rr, 6 + j).font = F_INPUT
    ws_ref.cell(rr, 9).font = F_BASE_B
    ws_ref.cell(rr, 9).fill = FILL_ITEM
    ws_ref.cell(rr, 10).font = Font(name="Calibri", size=8, color="595959")

def qty_ref(item_idx, well):
    col = 6 + WELLS.index(well)
    return f"Справочник!${get_column_letter(col)}${NOM_R0 + item_idx}"

def uom_ref(item_idx):
    return f"Справочник!$C${NOM_R0 + item_idx}"

ws_ref.freeze_panes = "A5"


# ============================================================================
# РАСЧЕТ МОДЕЛИ В PYTHON (для значений-констант и примечаний)
# ============================================================================
OP_DATE = {(w, o): d for w, o, d in OPERATIONS}

MITIGATION = {
    "Труба обсадная ОК-426": "СРЫВ {n} сут. Мероприятие: закупка со склада готовой продукции завода "
                             "(исключает этап изготовления, -35 сут) либо перераспределение трубы с ЮГ-2/ЮГ-3 "
                             "с последующим восполнением.",
    "Труба обсадная ОК-324": "СРЫВ {n} сут. Мероприятие: закупка со склада готовой продукции завода "
                             "(-35 сут) либо ускоренная отправка отдельным вагоном без накопления партии (-6 сут).",
    "Тех.оснастка (ОК-426)": "СРЫВ {n} сут. Мероприятие: закупка складской позиции у дистрибьютора "
                             "(-25 сут) либо авиадоставка Владивосток — Петропавловск-Камчатский (-7 сут, малый вес/объем).",
}

model = {}   # (well, item_idx) -> dict
for well in WELLS:
    for idx, (name, uom, cat, op, qty, note) in enumerate(NOMENCLATURE):
        total = sum(CAT_DUR[cat])
        need = OP_DATE[(well, op)] - dt.timedelta(days=NEED_LEAD)
        req_start = need - dt.timedelta(days=total)
        planned = max(PROCUREMENT_READY, req_start - dt.timedelta(days=TARGET_BUFFER))
        forecast = planned + dt.timedelta(days=total)
        slack = (need - forecast).days
        status = "СРЫВ ГРАФИКА" if slack < 0 else ("РИСК" if slack < RISK_THRESHOLD else "В ГРАФИКЕ")
        if status == "СРЫВ ГРАФИКА":
            act = MITIGATION.get(name, "СРЫВ {n} сут. Требуется ускоренная схема поставки или сдвиг срока крепления.")
            act = act.format(n=abs(slack))
        elif status == "РИСК":
            act = f"Запас {slack} сут — ниже норматива {RISK_THRESHOLD} сут. Заказ разместить первым приоритетом, взять на еженедельный контроль."
        else:
            act = ""
        model[(well, idx)] = dict(name=name, uom=uom, cat=cat, op=op, qty=qty, total=total,
                                  need=need, req_start=req_start, planned=planned,
                                  forecast=forecast, slack=slack, status=status, action=act)


# ============================================================================
# ЛИСТЫ 3-5 — ГРАФИКИ ПО СКВАЖИНАМ
# ============================================================================
WCOLS = [("A", 5), ("B", 58), ("C", 9), ("D", 10), ("E", 9), ("F", 12), ("G", 12),
         ("H", 13), ("I", 13), ("J", 14), ("K", 13), ("L", 10), ("M", 15), ("N", 52)]
WHDR = ["№", "Наименование МТР / этапа закупки и доставки", "Ед. изм.", "Кол-во",
        "Длит., сут", "Начало этапа", "Окончание этапа", "ДАТА ПОТРЕБНОСТИ",
        "Требуемая дата старта", "Дата начала (план/факт)", "Прогноз поставки на скважину",
        "Запас, сут", "СТАТУС", "Примечание / компенсирующие мероприятия"]

well_item_row = {}   # (well, idx) -> строка заголовка позиции на листе скважины

for well in WELLS:
    ws = wb.create_sheet(well)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.outlinePr.summaryBelow = False
    for col, w in WCOLS:
        ws.column_dimensions[col].width = w

    ws.cell(1, 1, f"ГРАФИК ЗАКУПКИ И ДОСТАВКИ МТР ДЛЯ КРЕПЛЕНИЯ СКВАЖИНЫ {well}").font = F_TITLE
    ws.merge_cells("A1:N1")
    ws.cell(2, 1, f"Дата формирования {TODAY.strftime('%d.%m.%Y')}. "
                  f"Даты потребности и длительности этапов — с листа «Справочник». "
                  f"Желтые ячейки (столбец J) — ввод фактической даты размещения заказа. "
                  f"Группировка слева сворачивает этапы.").font = F_SUB
    ws.merge_cells("A2:N2")

    for i, h in enumerate(WHDR, start=1):
        ws.cell(4, i, h)
    ws.row_dimensions[4].height = 46
    style_range(ws, "A4:N4", font=F_HDR, fill=FILL_HDR, border=B_ALL, align=A_C)

    rr = 5
    for idx, (name, uom, cat, op, qty, note) in enumerate(NOMENCLATURE):
        m = model[(well, idx)]
        hdr = rr
        well_item_row[(well, idx)] = hdr
        s0, s1 = hdr + 1, hdr + len(STAGES)

        ws.cell(hdr, 1, idx + 1)
        ws.cell(hdr, 2, name)
        ws.cell(hdr, 3, f"={uom_ref(idx)}")
        ws.cell(hdr, 4, f"={qty_ref(idx, well)}")
        ws.cell(hdr, 5, f"=SUM(E{s0}:E{s1})")
        ws.cell(hdr, 6, f"Этап крепления: {op}")
        ws.merge_cells(start_row=hdr, start_column=6, end_row=hdr, end_column=7)
        ws.cell(hdr, 8, f"={need_ref(well, op)}")
        ws.cell(hdr, 9, f"=H{hdr}-E{hdr}")
        ws.cell(hdr, 10, m["planned"])
        ws.cell(hdr, 11, f"=G{s1}")
        ws.cell(hdr, 12, f"=H{hdr}-K{hdr}")
        ws.cell(hdr, 13, f'=IF(L{hdr}<0,"СРЫВ ГРАФИКА",IF(L{hdr}<{RISK_THRESHOLD},"РИСК","В ГРАФИКЕ"))')
        ws.cell(hdr, 14, m["action"])

        ws.row_dimensions[hdr].height = 30
        style_range(ws, f"A{hdr}:N{hdr}", font=F_ITEM, fill=FILL_ITEM, border=B_ITEM, align=A_C)
        ws.cell(hdr, 2).alignment = A_L
        ws.cell(hdr, 6).alignment = A_C
        ws.cell(hdr, 6).font = Font(name="Calibri", size=8, italic=True, color="595959")
        for cc in (8, 9, 10, 11):
            ws.cell(hdr, cc).number_format = DATE_FMT
        ws.cell(hdr, 10).fill = FILL_INPUT
        ws.cell(hdr, 10).font = F_INPUT
        ws.cell(hdr, 14).font = Font(name="Calibri", size=8, color="9C0006")
        ws.cell(hdr, 14).alignment = A_LT

        for k, st in enumerate(STAGES):
            r_ = s0 + k
            ws.cell(r_, 1, k + 1)
            ws.cell(r_, 2, f"=Справочник!$B${STAGE_R0 + k}")
            ws.cell(r_, 5, f"={dur_ref(cat, k)}")
            ws.cell(r_, 6, f"=J{hdr}" if k == 0 else f"=G{r_-1}")
            ws.cell(r_, 7, f"=F{r_}+E{r_}")
            ws.row_dimensions[r_].height = 24
            ws.row_dimensions[r_].outlineLevel = 1
            fill = FILL_BAND if k % 2 == 0 else None
            style_range(ws, f"A{r_}:N{r_}", font=F_BASE, fill=fill, border=B_ALL, align=A_C)
            ws.cell(r_, 2).alignment = A_L
            ws.cell(r_, 6).number_format = DATE_FMT
            ws.cell(r_, 7).number_format = DATE_FMT

        rr = s1 + 2

    last = rr - 2
    # условное форматирование статуса / запаса
    for st, style in CLR.items():
        ds = DifferentialStyle(font=Font(bold=True, color=style["text"]),
                               fill=PatternFill("solid", bgColor=style["cell"]))
        ws.conditional_formatting.add(
            f"L5:M{last}",
            Rule(type="expression", formula=[f'$M5="{st}"'], dxf=ds, stopIfTrue=True))

    ws.freeze_panes = "C5"
    ws.print_title_rows = "4:4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# ============================================================================
# ЛИСТ 6 — СВОДНЫЙ ГРАФИК (диаграмма Ганта)
# ============================================================================
GANTT_FROM = dt.date(2026, 9, 1)
GANTT_TO = dt.date(2027, 6, 30)
NDAYS = (GANTT_TO - GANTT_FROM).days + 1
T0 = 13                      # первая колонка шкалы времени (M)
TLAST = T0 + NDAYS - 1
TL0, TLL = get_column_letter(T0), get_column_letter(TLAST)

gs = wb.create_sheet("СВОДНЫЙ ГРАФИК")
gs.sheet_view.showGridLines = False

GCOLS = [(1, 5), (2, 9), (3, 44), (4, 8), (5, 9), (6, 11), (7, 11), (8, 11),
         (9, 11), (10, 9), (11, 15), (12, 46)]
for ci, w in GCOLS:
    gs.column_dimensions[get_column_letter(ci)].width = w
for ci in range(T0, TLAST + 1):
    gs.column_dimensions[get_column_letter(ci)].width = 2.3

gs.cell(1, 1, "СВОДНЫЙ ГРАФИК ЗАКУПКИ И ДОСТАВКИ МТР ДЛЯ КРЕПЛЕНИЯ СКВАЖИН ЮГ-1, ЮГ-2, ЮГ-3").font = F_TITLE
gs.merge_cells("A1:L1")
gs.cell(2, 1, f"Дата формирования {TODAY.strftime('%d.%m.%Y')}. "
              "Вертикальная линия — текущая дата (пересчитывается функцией СЕГОДНЯ при каждом открытии файла). "
              "Цвет линии в строке: ЗЕЛЕНАЯ — в графике, ЖЕЛТАЯ — риск, КРАСНАЯ — срыв графика.").font = F_SUB
gs.merge_cells("A2:J2")

# легенда
leg = [("Производственный цикл (заказ, изготовление, затарка)", "E2EFDA", "A9D08E", "1F4E79"),
       ("Транспортное плечо (ж/д, море, доставка на скважину)", "A9D08E", "A9D08E", "1F4E79")]
gs.cell(3, 1, "ЛЕГЕНДА:").font = F_BASE_B
lc = 2
for txt, f1, f2, _ in leg:
    c = gs.cell(3, lc); c.fill = PatternFill("solid", fgColor=f1); c.border = B_ALL
    gs.cell(3, lc + 1, txt).font = F_BASE
    gs.merge_cells(start_row=3, start_column=lc + 1, end_row=3, end_column=lc + 2)
    lc += 4
c = gs.cell(3, lc); c.fill = PatternFill("solid", fgColor="002060"); c.border = B_ALL
gs.cell(3, lc + 1, "Дата потребности МТР на скважине").font = F_BASE
gs.merge_cells(start_row=3, start_column=lc + 1, end_row=3, end_column=lc + 2)

# общий статус программы (худший по всем позициям)
gs.cell(2, 11, "ОБЩИЙ СТАТУС:").font = F_BASE_B
gs.cell(2, 11).alignment = Alignment(horizontal="right", vertical="center")
gs.cell(2, 12, '=IF(COUNTIF(K6:K50,"СРЫВ ГРАФИКА")>0,"СРЫВ ГРАФИКА",'
               'IF(COUNTIF(K6:K50,"РИСК")>0,"РИСК","В ГРАФИКЕ"))')
gs.cell(2, 12).alignment = A_C
gs.cell(2, 12).border = B_ALL

# --- шапка таблицы
GHDR = ["№", "Скважина", "Наименование МТР", "Ед. изм.", "Кол-во", "Дата начала (план/факт)",
        "Окончание изготовления", "Прогноз поставки", "ДАТА ПОТРЕБНОСТИ", "Запас, сут",
        "СТАТУС", "Примечание / компенсирующие мероприятия"]
bn = gs.cell(4, 1, "ПОЗИЦИИ ЗАКУПКИ И ДОСТАВКИ МТР")
gs.merge_cells(start_row=4, start_column=1, end_row=4, end_column=12)
for i, h in enumerate(GHDR, start=1):
    gs.cell(5, i, h)
gs.row_dimensions[4].height = 18
gs.row_dimensions[5].height = 46
style_range(gs, "A4:L5", font=F_HDR, fill=FILL_HDR, border=B_ALL, align=A_C)
bn.alignment = A_C

# --- шкала времени: строка 4 месяцы, строка 5 числа
MON = ["янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек"]
d = GANTT_FROM
ci = T0
month_start = ci
while d <= GANTT_TO:
    cell = gs.cell(5, ci, d)
    cell.number_format = "d"
    cell.font = Font(name="Calibri", size=6, color="595959")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = Border(left=Side(style="hair", color="D9D9D9"),
                         right=Side(style="hair", color="D9D9D9"),
                         top=_thin, bottom=_thin)
    nxt = d + dt.timedelta(days=1)
    if nxt.month != d.month or nxt > GANTT_TO:
        gs.merge_cells(start_row=4, start_column=month_start, end_row=4, end_column=ci)
        mc = gs.cell(4, month_start, f"{MON[d.month-1]} {d.year}")
        mc.font = Font(name="Calibri", size=8, bold=True, color="FFFFFF")
        mc.fill = FILL_HDR
        mc.alignment = A_C
        mc.border = B_ALL
        month_start = ci + 1
    d = nxt
    ci += 1

# --- строки данных
GROW0 = 6
rr = GROW0
gantt_rows = {}
well_group_rows = []
n = 0
for well in WELLS:
    gc = gs.cell(rr, 1, f"СКВАЖИНА {well}")
    gs.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=12)
    style_range(gs, f"A{rr}:L{rr}", font=Font(name="Calibri", size=11, bold=True, color="FFFFFF"),
                fill=FILL_SUB, border=B_ALL, align=A_L)
    for ci2 in range(T0, TLAST + 1):
        gs.cell(rr, ci2).fill = PatternFill("solid", fgColor="DDEBF7")
    gs.row_dimensions[rr].height = 20
    well_group_rows.append(rr)
    rr += 1
    for idx in range(len(NOMENCLATURE)):
        n += 1
        src = well_item_row[(well, idx)]
        gs.cell(rr, 1, n)
        gs.cell(rr, 2, well)
        gs.cell(rr, 3, f"='{well}'!B{src}")
        gs.cell(rr, 4, f"='{well}'!C{src}")
        gs.cell(rr, 5, f"='{well}'!D{src}")
        gs.cell(rr, 6, f"='{well}'!J{src}")
        gs.cell(rr, 7, f"='{well}'!G{src+3}")      # окончание 3-го этапа = конец произв. цикла
        gs.cell(rr, 8, f"='{well}'!K{src}")
        gs.cell(rr, 9, f"='{well}'!H{src}")
        gs.cell(rr, 10, f"='{well}'!L{src}")
        gs.cell(rr, 11, f"='{well}'!M{src}")
        gs.cell(rr, 12, f"='{well}'!N{src}")
        gs.row_dimensions[rr].height = 17
        style_range(gs, f"A{rr}:L{rr}", font=F_BASE, border=B_ALL, align=A_C)
        gs.cell(rr, 3).alignment = A_L
        gs.cell(rr, 3).font = F_BASE_B
        gs.cell(rr, 12).alignment = A_L
        gs.cell(rr, 12).font = Font(name="Calibri", size=8, color="9C0006")
        for cc in (6, 7, 8, 9):
            gs.cell(rr, cc).number_format = DATE_FMT
        gs.cell(rr, 9).font = F_BASE_B
        for ci2 in range(T0, TLAST + 1):
            gs.cell(rr, ci2).border = Border(left=Side(style="hair", color="EDEDED"),
                                             right=Side(style="hair", color="EDEDED"),
                                             top=Side(style="hair", color="EDEDED"),
                                             bottom=Side(style="hair", color="EDEDED"))
        gantt_rows[(well, idx)] = rr
        rr += 1
GLAST = rr - 1

# --- условное форматирование: статусные ячейки
for st, style in CLR.items():
    ds = DifferentialStyle(font=Font(bold=True, color=style["text"]),
                           fill=PatternFill("solid", bgColor=style["cell"]))
    gs.conditional_formatting.add(
        f"J{GROW0}:K{GLAST}",
        Rule(type="expression", formula=[f'$K{GROW0}="{st}"'], dxf=ds, stopIfTrue=True))

for st, style in CLR.items():
    ds = DifferentialStyle(font=Font(bold=True, size=11, color=style["text"]),
                           fill=PatternFill("solid", bgColor=style["cell"]))
    gs.conditional_formatting.add("L2", Rule(type="expression", formula=[f'$L$2="{st}"'],
                                             dxf=ds, stopIfTrue=True))

# --- условное форматирование: диаграмма Ганта
AREA = f"{TL0}{GROW0}:{TLL}{GLAST}"
R = GROW0
TC = f"{TL0}$5"          # ссылка на дату в шапке шкалы

def add(formula, fill_color, stop=True, font=None):
    ds = DifferentialStyle(fill=PatternFill("solid", bgColor=fill_color), font=font)
    gs.conditional_formatting.add(AREA, Rule(type="expression", formula=[formula],
                                             dxf=ds, stopIfTrue=stop))

# 1-4: вертикальная линия текущей даты, окрашенная по статусу строки
for st in ("СРЫВ ГРАФИКА", "РИСК", "В ГРАФИКЕ"):
    add(f'AND({TC}=TODAY(),$K{R}="{st}")', CLR[st]["line"])
add(f'{TC}=TODAY()', "404040")
# 5: маркер даты потребности
add(f'AND($I{R}<>"",{TC}=$I{R})', "002060")
# 6-8: транспортное плечо
for st in ("СРЫВ ГРАФИКА", "РИСК", "В ГРАФИКЕ"):
    add(f'AND($G{R}<>"",{TC}>$G{R},{TC}<=$H{R},$K{R}="{st}")', CLR[st]["transport"])
# 9-11: производственный цикл
for st in ("СРЫВ ГРАФИКА", "РИСК", "В ГРАФИКЕ"):
    add(f'AND($F{R}<>"",{TC}>=$F{R},{TC}<=$G{R},$K{R}="{st}")', CLR[st]["prod"])
# 12: выходные дни
add(f'WEEKDAY({TC},2)>5', "F2F2F2", stop=False)

ds_hdr = DifferentialStyle(fill=PatternFill("solid", bgColor="404040"),
                           font=Font(bold=True, color="FFFFFF"))
gs.conditional_formatting.add(f"{TL0}5:{TLL}5",
                              Rule(type="expression", formula=[f"{TL0}$5=TODAY()"],
                                   dxf=ds_hdr, stopIfTrue=True))

gs.freeze_panes = f"{TL0}{GROW0}"
gs.auto_filter.ref = f"A5:L{GLAST}"
gs.print_title_rows = "4:5"
gs.page_setup.orientation = "landscape"
gs.page_setup.fitToWidth = 1
gs.page_setup.fitToHeight = 0
gs.sheet_properties.pageSetUpPr.fitToPage = True


# ============================================================================
# ЛИСТ 7 — СВОДНАЯ ПОТРЕБНОСТЬ
# ============================================================================
ss = wb.create_sheet("Сводная потребность")
ss.sheet_view.showGridLines = False
for ci, w in [(1, 5), (2, 46), (3, 10), (4, 11), (5, 11), (6, 11), (7, 13), (8, 15),
              (9, 14), (10, 12), (11, 16)]:
    ss.column_dimensions[get_column_letter(ci)].width = w

ss.cell(1, 1, "СВОДНАЯ ПОТРЕБНОСТЬ В МТР НА КРЕПЛЕНИЕ СКВАЖИН ЮГ-1, ЮГ-2, ЮГ-3").font = F_TITLE
ss.merge_cells("A1:K1")
ss.cell(2, 1, "Объем закупки по номенклатуре в целом по программе. Статус — худший из трех скважин.").font = F_SUB
ss.merge_cells("A2:K2")

SHDR = ["№", "Наименование МТР", "Ед. изм.", "ЮГ-1", "ЮГ-2", "ЮГ-3", "ИТОГО к закупке",
        "Цикл поставки, сут", "Самая ранняя потребность", "Мин. запас, сут", "ХУДШИЙ СТАТУС"]
for i, h in enumerate(SHDR, start=1):
    ss.cell(4, i, h)
ss.row_dimensions[4].height = 42
style_range(ss, "A4:K4", font=F_HDR, fill=FILL_HDR, border=B_ALL, align=A_C)

sr = 5
for idx, (name, uom, cat, op, qty, note) in enumerate(NOMENCLATURE):
    refs_need = [f"'{w}'!H{well_item_row[(w, idx)]}" for w in WELLS]
    refs_slack = [f"'{w}'!L{well_item_row[(w, idx)]}" for w in WELLS]
    ss.cell(sr, 1, idx + 1)
    ss.cell(sr, 2, f"=Справочник!B{NOM_R0 + idx}")
    ss.cell(sr, 3, f"=Справочник!C{NOM_R0 + idx}")
    for j, w in enumerate(WELLS):
        ss.cell(sr, 4 + j, f"='{w}'!D{well_item_row[(w, idx)]}")
    ss.cell(sr, 7, f"=SUM(D{sr}:F{sr})")
    ss.cell(sr, 8, f"='{WELLS[0]}'!E{well_item_row[(WELLS[0], idx)]}")
    ss.cell(sr, 9, f"=MIN({','.join(refs_need)})")
    ss.cell(sr, 10, f"=MIN({','.join(refs_slack)})")
    ss.cell(sr, 11, f'=IF(J{sr}<0,"СРЫВ ГРАФИКА",IF(J{sr}<{RISK_THRESHOLD},"РИСК","В ГРАФИКЕ"))')
    ss.row_dimensions[sr].height = 22
    style_range(ss, f"A{sr}:K{sr}", font=F_BASE, border=B_ALL, align=A_C)
    ss.cell(sr, 2).alignment = A_L
    ss.cell(sr, 2).font = F_BASE_B
    ss.cell(sr, 7).font = F_BASE_B
    ss.cell(sr, 7).fill = FILL_ITEM
    ss.cell(sr, 9).number_format = DATE_FMT
    ss.cell(sr, 9).font = F_BASE_B
    sr += 1

for st, style in CLR.items():
    ds = DifferentialStyle(font=Font(bold=True, color=style["text"]),
                           fill=PatternFill("solid", bgColor=style["cell"]))
    ss.conditional_formatting.add(
        f"J5:K{sr-1}",
        Rule(type="expression", formula=[f'$K5="{st}"'], dxf=ds, stopIfTrue=True))

ss.freeze_panes = "A5"
ss.auto_filter.ref = f"A4:K{sr-1}"

# ============================================================================
OUT = "/home/user/nestro/mtr_schedule/График_закупки_и_доставки_МТР_крепление_ЮГ-1_ЮГ-2_ЮГ-3.xlsx"
wb.active = wb.sheetnames.index("СВОДНЫЙ ГРАФИК")
wb.save(OUT)
print("OK ->", OUT)
print(f"позиций: {len(NOMENCLATURE)} x {len(WELLS)} = {len(NOMENCLATURE)*len(WELLS)}; "
      f"шкала {GANTT_FROM:%d.%m.%Y}-{GANTT_TO:%d.%m.%Y} ({NDAYS} дн.), строки {GROW0}-{GLAST}")
by_st = {}
for k, v in model.items():
    by_st[v["status"]] = by_st.get(v["status"], 0) + 1
print("статусы:", by_st)
for w in WELLS:
    bad = [(model[(w, i)]["name"], model[(w, i)]["slack"]) for i in range(len(NOMENCLATURE))
           if model[(w, i)]["status"] != "В ГРАФИКЕ"]
    print(f"  {w}: " + ("; ".join(f"{n} ({s:+d})" for n, s in bad) if bad else "все в графике"))
